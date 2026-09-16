"""Lectura e interpretación del Excel del levantamiento inicial.

Solo transforma datos; no toca la base. La escritura está en import_excel.py.
Las reglas de interpretación son las aprobadas el 2026-09-14 (docs/flujos.md, sección 0).
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

HOJAS = {
    "VEX": "VEX",
    "FTC": "FTC",
    "Herramientas": "HERRAMIENTAS",
    "Herramientas eléctricas": "HERRAMIENTAS_ELECTRICAS",
    "Consumibles": "CONSUMIBLES",
    "Común": "COMUN",
    "Sin clasificar": "SIN_CLASIFICAR",
}

ENCABEZADO = [
    "N.º", "Ref. original", "Artículo", "Marca / Modelo", "Cantidad contada", "Estado",
    "Subcategoría", "Ubicación (llenar en sitio)", "Observaciones", "Pendiente / Por revisar",
    "Verificado en sitio",
]
FILA_ENCABEZADO = 4

# Subcategorías del Excel que en realidad eran estados; se convierten en pendientes.
SUBCATEGORIAS_QUE_SON_ESTADO = {"Cajas vacías / verificar", "Pendientes"}

# Unidades que son un envase: si el artículo no es consumible, hay que contar lo de adentro.
ENVASES = {"caja", "bolsa", "estuche", "contenedor", "canasta"}

# (texto normalizado con el que empieza la etiqueta, tipo de tarea)
ETIQUETAS_PENDIENTE = [
    ("contar fisicamente", "CONTAR"),
    ("verificar dato", "VERIFICAR_DATO"),
    ("abrir y revisar", "ABRIR_REVISAR"),
    ("registrar serie", "REGISTRAR_SERIE"),
    ("identificar", "IDENTIFICAR_ETIQUETAR"),
    ("falta pieza", "FALTA_PIEZA"),
    ("definir si es vex o ftc", "DEFINIR_VEX_FTC"),
    ("confirmar si esta vacio", "CONFIRMAR_VACIO"),
]

RE_RESGUARDO = re.compile(r"\bP13/\d{4}\b")
RE_SERIE = re.compile(r"[Ss]erie\s*:?\s*([A-Z0-9][A-Z0-9-]{5,})")


def normalizar(texto: str) -> str:
    """Minúsculas, sin acentos y sin espacios repetidos."""
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return " ".join(sin_acentos.lower().split())


def texto(valor) -> str:
    return "" if valor is None else " ".join(str(valor).split())


# ---------------------------------------------------------------------------
# Cantidades
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Cantidad:
    valor: int | None
    unidad: str
    estimada: bool
    desconocida: bool


def singular(palabra: str) -> str:
    if palabra.endswith("ciones"):
        return palabra[:-6] + "ción"
    if palabra.endswith("s") and len(palabra) > 3:
        return palabra[:-1]
    return palabra


def unidad_por_nombre(nombre: str) -> str:
    primera = singular(normalizar(nombre).split(" ", 1)[0]) if nombre else ""
    return {"juego": "juego", "kit": "kit", "pack": "paquete", "caja": "caja"}.get(primera, "pieza")


def interpretar_cantidad(valor, nombre_articulo: str = "") -> Cantidad:
    """Convierte el texto de "Cantidad contada" en número, unidad y marcas.

    Reglas aprobadas:
      "4 cajas"      -> 4 caja, exacta (se contó el envase)
      "~14"          -> 14, estimada
      "1 (mínimo)"   -> 1, estimada;  "2 o más" -> 2, estimada
      "bolsa, ~10"   -> 10 piezas, estimada
      "varias", "1 contenedor lleno" -> sin conteo (no se presta hasta contarlo)
    """
    t = texto(valor)
    por_nombre = unidad_por_nombre(nombre_articulo)
    if not t:
        return Cantidad(None, por_nombre, True, True)
    if re.fullmatch(r"\d+", t):
        return Cantidad(int(t), por_nombre, False, False)

    bajo = t.lower()
    if normalizar(t) in {"varias", "varios"} or "lleno" in bajo:
        return Cantidad(None, por_nombre, True, True)

    estimada = "~" in t or "mínimo" in bajo or "minimo" in bajo or re.search(r"\bo m[aá]s\b", bajo) is not None

    envase_y_piezas = re.fullmatch(r"([a-záéíóúñ ]+),\s*~\s*(\d+)", bajo)
    if envase_y_piezas:
        return Cantidad(int(envase_y_piezas[2]), "pieza", True, False)

    m = re.match(r"~?\s*(\d+)\s*(.*)", bajo)
    if not m:
        return Cantidad(None, por_nombre, True, True)
    resto = re.sub(r"\(.*?\)", "", m[2])
    resto = re.sub(r"\bo m[aá]s\b", "", resto).strip(" ,")
    unidad = singular(resto.split()[0]) if resto else por_nombre
    return Cantidad(int(m[1]), unidad, estimada, False)


# ---------------------------------------------------------------------------
# Estado físico
# ---------------------------------------------------------------------------
def interpretar_estado(valor) -> str | None:
    n = normalizar(texto(valor))
    if not n:
        return None
    # Solo "VACÍO" a secas. "Usados, mayoría vacíos" (organizadores) es uso normal, no un faltante.
    if n.startswith("vacio"):
        return "VACIO"
    if "incomplet" in n:
        return "INCOMPLETO"
    if "danad" in n or "roto" in n:
        return "DANADO"
    if any(p in n for p in ("sellad", "emplayad", "cerrad")):
        return "SIN_ABRIR"
    if "nuev" in n:
        return "NUEVO"
    if any(p in n for p in ("usad", "parcial", "operativ", "funcion", "en uso", "instalad",
                            "completo", "montad", "abiert", "oxid")):
        return "USADO"
    return None


# ---------------------------------------------------------------------------
# Renglones del Excel
# ---------------------------------------------------------------------------
@dataclass
class Tarea:
    tipo: str
    descripcion: str
    prioridad: int = 0


@dataclass
class ArticuloExcel:
    hoja: str
    fila: int
    ref_foto: int
    nombre: str
    marca_modelo: str | None
    categoria: str
    subcategoria: str | None
    cantidad: Cantidad
    cantidad_texto: str
    estado_fisico: str | None
    estado_fisico_texto: str | None
    ubicacion: str | None
    observaciones: str | None
    num_resguardo: str | None
    num_serie: str | None
    es_consumible: bool
    etiquetado: str
    estado_inventario: str
    tareas: list[Tarea] = field(default_factory=list)
    datos_origen: dict = field(default_factory=dict)
    advertencias: list[str] = field(default_factory=list)


def tareas_de_etiquetas(pendiente: str, observaciones: str | None, advertencias: list[str]) -> list[Tarea]:
    tareas = []
    for etiqueta in (e.strip() for e in pendiente.split(";")):
        if not etiqueta:
            continue
        n = normalizar(etiqueta)
        tipo = next((t for inicio, t in ETIQUETAS_PENDIENTE if n.startswith(inicio)), None)
        if tipo is None:
            advertencias.append(f'Etiqueta de pendiente no reconocida: "{etiqueta}" (se registró como "Verificar dato")')
            tipo = "VERIFICAR_DATO"
        descripcion = etiqueta if not observaciones else f"{etiqueta}. {observaciones}"
        tareas.append(Tarea(tipo, descripcion))
    return tareas


def agregar_tarea(tareas: list[Tarea], nueva: Tarea) -> None:
    """Una tarea por tipo: si ya existe, se suma la descripción."""
    for t in tareas:
        if t.tipo == nueva.tipo:
            if nueva.descripcion not in t.descripcion:
                t.descripcion = f"{t.descripcion} · {nueva.descripcion}"
            t.prioridad = max(t.prioridad, nueva.prioridad)
            return
    tareas.append(nueva)


def convertir_fila(hoja: str, fila: int, valores: list) -> ArticuloExcel:
    (_, ref, nombre, marca, cant, estado, subcat, ubic, obs, pendiente, verificado) = (valores + [None] * 11)[:11]
    advertencias: list[str] = []
    categoria = HOJAS[hoja]
    nombre = texto(nombre)
    marca = texto(marca) or None
    obs = texto(obs) or None
    subcat = texto(subcat) or None
    cantidad_texto = texto(cant)
    cantidad = interpretar_cantidad(cant, nombre)
    estado_fisico = interpretar_estado(estado)
    es_consumible = categoria == "CONSUMIBLES"

    busqueda = " ".join(filter(None, [nombre, marca, obs]))
    resguardo = RE_RESGUARDO.search(busqueda)
    serie = next((s for s in RE_SERIE.findall(obs or "") if re.search(r"\d", s)), None)

    tareas = tareas_de_etiquetas(texto(pendiente), obs, advertencias)
    if subcat in SUBCATEGORIAS_QUE_SON_ESTADO:
        if subcat == "Cajas vacías / verificar":
            agregar_tarea(tareas, Tarea("CONFIRMAR_VACIO", "Registrado como caja vacía o por verificar en el levantamiento"))
        subcat = None
    if estado_fisico == "VACIO":
        agregar_tarea(tareas, Tarea(
            "LOCALIZAR_CONTENIDO", "Se recibió vacío: localizar su contenido o darlo por faltante", prioridad=1))
    if cantidad.desconocida or cantidad.estimada:
        agregar_tarea(tareas, Tarea("CONTAR", f'Cantidad registrada como "{cantidad_texto}": contar físicamente'))
    elif cantidad.unidad in ENVASES and not es_consumible:
        agregar_tarea(tareas, Tarea("CONTAR", f"Contar lo que hay dentro ({cantidad_texto})"))

    tipos = {t.tipo for t in tareas}
    if categoria == "HERRAMIENTAS_ELECTRICAS" or resguardo or serie or "REGISTRAR_SERIE" in tipos:
        etiquetado = "INDIVIDUAL"
    elif es_consumible:
        etiquetado = "LOTE"
    else:
        etiquetado = "CONTENEDOR"

    if categoria == "SIN_CLASIFICAR":
        estado_inventario = "SIN_CLASIFICAR"
    elif cantidad.estimada or cantidad.desconocida:
        estado_inventario = "POR_CONTAR"
    else:
        estado_inventario = "POR_VERIFICAR"

    if not isinstance(ref, int):
        raise ValueError(f'{hoja}, renglón {fila}: "Ref. original" debe ser un número entero, llegó {ref!r}')
    if not nombre:
        raise ValueError(f"{hoja}, renglón {fila}: falta el nombre del artículo")

    return ArticuloExcel(
        hoja=hoja, fila=fila, ref_foto=ref, nombre=nombre, marca_modelo=marca, categoria=categoria,
        subcategoria=subcat, cantidad=cantidad, cantidad_texto=cantidad_texto,
        estado_fisico=estado_fisico, estado_fisico_texto=texto(estado) or None,
        ubicacion=texto(ubic) or None, observaciones=obs,
        num_resguardo=resguardo[0] if resguardo else None, num_serie=serie,
        es_consumible=es_consumible, etiquetado=etiquetado, estado_inventario=estado_inventario,
        tareas=tareas, advertencias=advertencias,
        datos_origen={"hoja": hoja, "fila": fila,
                      **{col: (None if v is None else str(v)) for col, v in zip(ENCABEZADO, valores)}},
    )


def leer_excel(ruta: Path) -> list[ArticuloExcel]:
    libro = load_workbook(ruta, data_only=True, read_only=True)
    faltan = [h for h in HOJAS if h not in libro.sheetnames]
    if faltan:
        raise ValueError(f"Al Excel le faltan las hojas: {', '.join(faltan)}")

    articulos: list[ArticuloExcel] = []
    for hoja in HOJAS:
        filas = libro[hoja].iter_rows(min_row=FILA_ENCABEZADO, values_only=True)
        encabezado = [texto(v) for v in next(filas)][: len(ENCABEZADO)]
        if encabezado != ENCABEZADO:
            raise ValueError(f'La hoja "{hoja}" no tiene el encabezado esperado en la fila {FILA_ENCABEZADO}: {encabezado}')
        for numero, valores in enumerate(filas, start=FILA_ENCABEZADO + 1):
            valores = list(valores[: len(ENCABEZADO)])
            if all(v in (None, "") for v in valores):
                continue
            articulos.append(convertir_fila(hoja, numero, valores))
    libro.close()

    vistos: dict[int, ArticuloExcel] = {}
    for a in articulos:
        if a.ref_foto in vistos:
            otro = vistos[a.ref_foto]
            raise ValueError(f"Ref. original {a.ref_foto} repetida: {otro.hoja} renglón {otro.fila} y {a.hoja} renglón {a.fila}")
        vistos[a.ref_foto] = a
    return articulos
