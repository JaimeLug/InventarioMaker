"""Incorpora el levantamiento 3 (InventarioRobotica3erFaltaVerificarCant.xlsx). Es la última importación desde Excel:
a partir de aquí el inventario se lleva en la app.

Uso:
  python scripts/levantamiento3.py [archivo.xlsx] [--simular] [--nube]

Qué hace (aprobado en la Fase 4b, 2026-09-17):
  * Hoja VEX (pasó de 5 a 27 renglones, en tres ramas: V5, EDR y Campo):
      - Los 5 artículos VEX que ya existían conservan su código y se actualizan:
        A-0105 Classroom & Competition Base Kit, A-0128 Caja VEX EDR (alargada), A-0130 V5 Robot Kit (3 → 2 cajas,
        con ajuste), A-0131 Batería V5, A-0132 Robot V5 armado.
      - Lo que viene "por caja" del kit de aula no se da de alta suelto (se contaría dos veces): queda como
        contenido esperado del kit A-0105, en una plantilla, para desglosarlo con la app.
      - Lo "suelto" (1 Cortex, 1 VEXnet Joystick, 2 VEXnet Key) y el resto de renglones se dan de alta.
  * Las listas de contenido de fábrica (REV Starter Kit V3, goBILDA Strafer, REV Control & Power Bundle) se cargan
    como plantillas de kit. Sus columnas "encontrado" y notas NO se aplican: se revisan una por una en la app.
  * Las demás hojas no cambiaron respecto del primer Excel: no se tocan.

Es idempotente: correrlo dos veces no duplica nada.
"""
from __future__ import annotations

import argparse
import sys
import uuid
from pathlib import Path

from openpyxl import load_workbook
from psycopg.rows import tuple_row
from psycopg.types.json import Jsonb

import db
from inventario_excel import convertir_fila, texto

ARCHIVO = db.RAIZ / "InventarioRobotica3erFaltaVerificarCant.xlsx"
ESPACIO = uuid.UUID("0b8d2f7e-6a1c-4a35-8f5e-3c2d1e0a9b77")

# N.º del renglón en la hoja VEX nueva → código del artículo que ya existe.
EXISTENTES = {7: "A-0105", 22: "A-0128", 1: "A-0130", 5: "A-0131", 4: "A-0132"}

# Renglones que traen una parte suelta y otra dentro de cada caja del kit de aula.
SUELTOS = {8: "1", 9: "1", 10: "2"}

# Contenido de cada caja del kit de aula (A-0105): cantidad esperada por caja (None = "varios").
POR_CAJA = {8: 1, 9: 1, 10: 1, 11: 1, 12: None, 13: None, 14: 1, 15: 1, 16: 1, 17: 1, 18: None, 19: None, 20: None}

PLANTILLA_VEX = "VEX Classroom & Competition Base Kit (según levantamiento 3)"


def uid(*partes) -> uuid.UUID:
    return uuid.uuid5(ESPACIO, "/".join(str(p) for p in partes))


def filas_vex(libro) -> list[dict]:
    hoja = libro["VEX"]
    encabezado, filas = None, []
    for valores in hoja.iter_rows(values_only=True):
        if valores and valores[0] == "N.º":
            encabezado = [texto(v) for v in valores]
            continue
        if encabezado and valores and isinstance(valores[0], int):
            filas.append(dict(zip(encabezado, valores)))
    if len(filas) != 27:
        sys.exit(f"La hoja VEX debería tener 27 renglones y tiene {len(filas)}: revisa el archivo.")
    return filas


def valores_para_convertir(f: dict, cantidad: str | None = None, obs_extra: str | None = None) -> list:
    """Arma el renglón con el orden del primer Excel (con 'Ref. original' en la segunda columna)."""
    obs = " ".join(filter(None, [texto(f.get("Observaciones")), obs_extra])) or None
    return [f["N.º"], 900000 + f["N.º"], f["Artículo"], f.get("Marca / Modelo"), cantidad or f.get("Cantidad contada"),
            f.get("Estado"), f.get("Subcategoría"), f.get("Ubicación (llenar en sitio)"), obs,
            f.get("Pendiente / Por revisar"), f.get("Verificado en sitio")]


def plantillas_bom(libro) -> list[dict]:
    salida = []
    for hoja, nombre, sku in (("BOM REV Starter Kit V3", "REV FTC Starter Kit V3", "REV-45-1883"),
                              ("BOM goBILDA Strafer", "goBILDA Strafer Chassis Kit", "3209-0001-0005"),
                              ("BOM Control & Power Bundle", "REV Control & Power Bundle", "REV-35-1906")):
        ws = libro[hoja]
        titulo = texto(ws.cell(1, 1).value)
        nota = texto(ws.cell(2, 1).value)
        lineas, seccion, dentro = [], None, False
        for valores in ws.iter_rows(values_only=True):
            c0 = texto(valores[0]) if valores else ""
            if c0 in ("N.º parte", "SKU"):
                dentro = True
                continue
            if not dentro:
                continue
            if c0.upper().startswith("RESUMEN"):
                break
            if c0 and all(v in (None, "") for v in valores[1:]):
                seccion = c0.title()
                continue
            if c0 and isinstance(valores[2], (int, float)):
                descripcion = texto(valores[1])
                notas = texto(valores[6]) if len(valores) > 6 else ""
                lineas.append({
                    "seccion": seccion, "sku": c0, "descripcion": descripcion, "cantidad": int(valores[2]),
                    "unidad": "paquete" if "pack" in descripcion.lower() else "pieza",
                    # En Control & Power el gamepad no se entregó (se sustituyó por controles Logitech).
                    "nota": notas if "NO SE CONSIDERA FALTANTE" in notas.upper() else None,
                })
        salida.append({"nombre": nombre, "sku": sku, "categoria": "FTC", "fuente": f"{titulo} (hoja {hoja})", "nota": nota,
                       "lineas": lineas})
    return salida


def plantilla_vex(filas: list[dict]) -> dict:
    lineas = []
    for f in filas:
        if f["N.º"] in POR_CAJA:
            lineas.append({
                "seccion": "Contenido de cada caja", "sku": texto(f.get("Marca / Modelo")) or None, "descripcion": f["Artículo"],
                "cantidad": POR_CAJA[f["N.º"]], "unidad": "paquete" if "paquete" in texto(f.get("Cantidad contada")) else "pieza",
                "nota": f'Levantamiento 3: "{texto(f.get("Cantidad contada"))}". {texto(f.get("Observaciones"))}'.strip(),
            })
    return {"nombre": PLANTILLA_VEX, "sku": None, "categoria": "VEX",
            "fuente": "Hoja VEX del levantamiento 3 (renglones 8 a 20)",
            "nota": "Lo que el levantamiento registró dentro de las cajas de aula. No es lista de fábrica: confirmar al abrir.",
            "lineas": lineas}


def cargar_plantilla(cur, p: dict, informe: dict) -> None:
    if cur.execute("select 1 from public.plantilla_kit where nombre = %s", (p["nombre"],)).fetchone():
        return
    pid = uid("plantilla", p["nombre"])
    cur.execute("insert into public.plantilla_kit (id, nombre, sku, categoria, fuente, nota) values (%s, %s, %s, %s, %s, %s)",
                (pid, p["nombre"], p["sku"], p["categoria"], p["fuente"], p["nota"]))
    for orden, l in enumerate(p["lineas"], start=1):
        cur.execute("insert into public.plantilla_kit_linea (plantilla_id, orden, seccion, sku, descripcion, cantidad, unidad, nota) "
                    "values (%s, %s, %s, %s, %s, %s, %s, %s)",
                    (pid, orden, l["seccion"], l["sku"], l["descripcion"], l["cantidad"], l["unidad"], l["nota"]))
    informe["plantillas"].append(f'{p["nombre"]} ({len(p["lineas"])} renglones)')


def tareas(cur, articulo_id, a, informe: dict) -> None:
    for t in a.tareas:
        cur.execute("insert into public.tarea_pendiente (articulo_id, tipo, descripcion, origen, prioridad) values (%s, %s, %s, 'IMPORTACION', %s) "
                    "on conflict (articulo_id, tipo) where origen = 'IMPORTACION' do nothing",
                    (articulo_id, t.tipo, t.descripcion, t.prioridad))
        informe["tareas"] += cur.rowcount


def origen(f: dict) -> dict:
    return {"levantamiento": 3, "hoja": "VEX", "n": f["N.º"], **{k: (None if v is None else str(v)) for k, v in f.items() if k}}


def aplicar(conn, archivo: Path) -> dict:
    libro = load_workbook(archivo, data_only=True)
    filas = filas_vex(libro)
    informe = {"actualizados": [], "nuevos": [], "ajustes": [], "tareas": 0, "plantillas": []}
    with conn.cursor(row_factory=tuple_row) as cur:
        for f in filas:
            n = f["N.º"]
            if n in POR_CAJA and n not in SUELTOS:
                continue
            obs_extra = None
            if n in SUELTOS:
                obs_extra = f"Registrado el suelto; lo de cada caja está en el contenido del kit A-0105 (\"{texto(f.get('Cantidad contada'))}\")."
            a = convertir_fila("VEX", n, valores_para_convertir(f, SUELTOS.get(n), obs_extra))

            if n in EXISTENTES:
                codigo = EXISTENTES[n]
                fila = cur.execute("select id, cantidad, cantidad_texto from public.articulo where codigo = %s", (codigo,)).fetchone()
                if fila is None:
                    sys.exit(f"No encontré {codigo} en la base.")
                articulo_id, cantidad_actual, texto_actual = fila
                cur.execute("""update public.articulo set nombre = %s, marca_modelo = %s, subcategoria = %s, estado_fisico_texto = %s,
                                      observaciones = %s, cantidad_texto = %s, datos_origen = %s
                                where id = %s and (nombre, subcategoria, cantidad_texto) is distinct from (%s, %s, %s)""",
                            (a.nombre, a.marca_modelo, a.subcategoria, a.estado_fisico_texto, a.observaciones, a.cantidad_texto,
                             Jsonb(origen(f)), articulo_id, a.nombre, a.subcategoria, a.cantidad_texto))
                if cur.rowcount:
                    informe["actualizados"].append(f"{codigo} {a.nombre}")
                if a.cantidad.valor is not None and cantidad_actual is not None and a.cantidad.valor != cantidad_actual:
                    cur.execute("""insert into public.movimiento (articulo_id, tipo, cantidad, nota, origen, comando_id)
                                   values (%s, 'AJUSTE_CONTEO', %s, %s, 'IMPORTACION', %s) on conflict (comando_id) do nothing""",
                                (articulo_id, a.cantidad.valor - cantidad_actual,
                                 f'Levantamiento 3: "{a.cantidad_texto}" (antes "{texto_actual}")', uid("ajuste", n)))
                    if cur.rowcount:
                        informe["ajustes"].append(f"{codigo}: {cantidad_actual} → {a.cantidad.valor}")
                tareas(cur, articulo_id, a, informe)
                continue

            existe = cur.execute("select id, codigo from public.articulo where datos_origen ->> 'levantamiento' = '3' "
                                 "and (datos_origen ->> 'n')::int = %s", (n,)).fetchone()
            if existe:
                articulo_id = existe[0]
            else:
                articulo_id, codigo = cur.execute(
                    """insert into public.articulo (nombre, marca_modelo, categoria, subcategoria, cantidad_texto, cantidad_estimada,
                         conteo_desconocido, unidad, estado_inventario, estado_fisico, estado_fisico_texto, etiquetado, num_serie,
                         observaciones, es_consumible, datos_origen)
                       values (%s, %s, 'VEX', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, false, %s) returning id, codigo""",
                    (a.nombre, a.marca_modelo, a.subcategoria, a.cantidad_texto, a.cantidad.estimada, a.cantidad.desconocida,
                     a.cantidad.unidad if a.cantidad.unidad != "por" else "pieza", a.estado_inventario, a.estado_fisico,
                     a.estado_fisico_texto, a.etiquetado, a.num_serie, a.observaciones, Jsonb(origen(f)))).fetchone()
                informe["nuevos"].append(f"{codigo} {a.nombre} ({a.cantidad_texto})")
            if a.cantidad.valor and not a.cantidad.desconocida:
                cur.execute("""insert into public.movimiento (articulo_id, tipo, cantidad, nota, origen, comando_id)
                               values (%s, 'ALTA', %s, %s, 'IMPORTACION', %s) on conflict (comando_id) do nothing""",
                            (articulo_id, a.cantidad.valor, f'Alta desde el levantamiento 3 (hoja VEX, N.º {n}, cantidad "{a.cantidad_texto}")',
                             uid("alta", n)))
            tareas(cur, articulo_id, a, informe)

        # El kit de aula: su contenido por caja queda para desglosarlo con la app.
        kit = cur.execute("select id from public.articulo where codigo = 'A-0105'").fetchone()[0]
        cur.execute("""insert into public.tarea_pendiente (articulo_id, tipo, descripcion, origen, prioridad)
                       values (%s, 'ABRIR_REVISAR', %s, 'IMPORTACION', 1)
                       on conflict (articulo_id, tipo) where origen = 'IMPORTACION' do nothing""",
                    (kit, f'Abrir y desglosar las cajas con la lista "{PLANTILLA_VEX}"'))
        informe["tareas"] += cur.rowcount

        for p in [*plantillas_bom(libro), plantilla_vex(filas)]:
            cargar_plantilla(cur, p, informe)

        descuadres = cur.execute("select count(*) from app.v_descuadres").fetchone()[0]
        if descuadres:
            raise RuntimeError(f"Quedaron {descuadres} descuadres: no se guarda nada.")
    return informe


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description="Incorpora el levantamiento 3 (hoja VEX y listas de contenido de fábrica)")
    p.add_argument("archivo", nargs="?", default=str(ARCHIVO))
    p.add_argument("--simular", action="store_true", help="hace todo y al final deshace")
    p.add_argument("--nube", action="store_true", help="en Supabase (datos en .env)")
    args = p.parse_args()

    print(f"Destino: {db.describir(args.nube)}")
    with db.conectar(nube=args.nube) as conn:
        informe = aplicar(conn, Path(args.archivo))
        if args.simular:
            conn.rollback()
        else:
            conn.commit()
    print("SIMULACIÓN: no se guardó nada." if args.simular else "Guardado.")
    for titulo, clave in (("Artículos actualizados", "actualizados"), ("Artículos nuevos", "nuevos"),
                          ("Ajustes de cantidad", "ajustes"), ("Plantillas de kit", "plantillas")):
        print(f"{titulo}: {len(informe[clave])}")
        for x in informe[clave]:
            print(f"  {x}")
    print(f"Pendientes nuevos: {informe['tareas']}")


if __name__ == "__main__":
    main()
